#!/usr/bin/env python3
"""SABZ official 4-slide deck (16:9) + assets. Facts: judge_prep verified set only."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

INK   = RGBColor(0x0B, 0x12, 0x20)
CARD  = RGBColor(0x14, 0x1D, 0x33)
CARD2 = RGBColor(0x10, 0x28, 0x1C)
LINE  = RGBColor(0x24, 0x30, 0x4D)
GREEN = RGBColor(0x00, 0xE6, 0x76)
TEXT  = RGBColor(0xED, 0xF2, 0xF7)
MUTED = RGBColor(0x93, 0xA1, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED   = RGBColor(0xF8, 0x71, 0x71)

HEAD = "Space Grotesk"
BODY = "Inter"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = INK
    return s

def rect(s, x, y, w, h, fill=CARD, line=LINE, lw=1.0, round_=True, radius=0.055):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try: shp.adjustments[0] = radius
        except Exception: pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, size=14, color=TEXT, font=BODY, bold=False,
         align=PP_ALIGN.LEFT, leading=1.12, anchor=MSO_ANCHOR.TOP, space_after=0):
    """runs: str, or list of paragraphs; each paragraph is list of (txt, dict) run tuples."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str): runs = [[(runs, {})]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = leading
        if space_after: p.space_after = Pt(space_after)
        for txt, ov in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(ov.get("size", size)); r.font.bold = ov.get("bold", bold)
            r.font.name = ov.get("font", font); r.font.color.rgb = ov.get("color", color)
    return tb

def logo(s, x=0.55, y=0.34, w=2.05):
    s.shapes.add_picture("../logo.png", Inches(x), Inches(y), width=Inches(w))

def kicker_title(s, kicker, title, tsize=38, y=1.32):
    text(s, 0.55, y, 12.2, 0.35, [[(kicker, {"color": GREEN, "font": HEAD, "bold": True, "size": 13})]])
    text(s, 0.55, y + 0.36, 12.25, 0.95, [[(title, {"font": HEAD, "bold": True, "size": tsize})]])

def footer(s, n):
    text(s, 0.55, 7.14, 6.0, 0.25, [[("SABZ · AI-ADAPTIVE TRAFFIC SIGNALS", {"color": MUTED, "size": 9.5})]])
    text(s, 12.35, 7.14, 0.45, 0.25, [[(str(n), {"color": MUTED, "size": 9.5})]], align=PP_ALIGN.RIGHT)

def bullet(s, x, y, w, txt, size=15.5, leading=1.18, h=0.62):
    text(s, x, y, w, h, [[("•  ", {"color": GREEN, "bold": True, "size": size}),
                          (txt, {"size": size})]], leading=leading)

# ============================================================ SLIDE 1: PROBLEM
s = slide(); logo(s)
kicker_title(s, "THE PROBLEM", "Traffic lights that cannot see", tsize=40)
bullet(s, 0.55, 2.85, 7.55, "Lahore moves 4.1 million vehicles through signals that follow a clock, not the traffic.", h=0.75)
bullet(s, 0.55, 3.72, 7.55, "The average commute takes 34.8 minutes, much of it spent idling at red lights.", h=0.75)
bullet(s, 0.55, 4.59, 7.55, "Fixed timers run blind cycles: an empty road can hold green while a long queue burns time and fuel on the cross street.", h=1.0)
bullet(s, 0.55, 5.62, 7.55, "The signal never sees which road actually needs the green.", h=0.62)
# right visual: the blind cycle
rect(s, 8.5, 2.8, 4.28, 3.42)
text(s, 8.78, 3.0, 3.75, 0.3, [[("THE BLIND CYCLE", {"color": GREEN, "font": HEAD, "bold": True, "size": 12.5})]])
rect(s, 8.78, 3.48, 1.78, 0.62, fill=CARD2, line=GREEN, lw=1.2)
text(s, 8.78, 3.62, 1.78, 0.35, [[("NS GREEN · 30s", {"font": HEAD, "bold": True, "size": 11.5})]], align=PP_ALIGN.CENTER)
rect(s, 10.72, 3.48, 1.78, 0.62, fill=CARD, line=RED, lw=1.2)
text(s, 10.72, 3.62, 1.78, 0.35, [[("EW GREEN · 30s", {"font": HEAD, "bold": True, "size": 11.5})]], align=PP_ALIGN.CENTER)
text(s, 8.78, 4.3, 3.72, 0.65, [[("Then repeat. The same split at rush hour and at 3 AM, whatever the road looks like.", {"color": MUTED, "size": 12.5})]], leading=1.25)
rect(s, 8.78, 5.12, 3.72, 0.02, fill=LINE, line=None, round_=False)
text(s, 8.78, 5.28, 1.85, 0.75, [[("4.1M", {"color": GREEN, "font": HEAD, "bold": True, "size": 24})]])
text(s, 10.6, 5.28, 1.9, 0.75, [[("34.8", {"color": GREEN, "font": HEAD, "bold": True, "size": 24}),
                                 (" min", {"color": MUTED, "size": 12})]])
text(s, 8.78, 5.86, 1.85, 0.3, [[("vehicles in Lahore", {"color": MUTED, "size": 10.5})]])
text(s, 10.6, 5.86, 1.9, 0.3, [[("average commute", {"color": MUTED, "size": 10.5})]])
text(s, 0.55, 6.5, 12.2, 0.4, [[("Result: ", {"color": GREEN, "bold": True, "size": 15}),
                                ("green time is wasted where nobody is waiting, and missing where queues are longest.", {"size": 15})]])
footer(s, 1)

# ============================================================ SLIDE 2: SOLUTION
s = slide(); logo(s)
kicker_title(s, "THE SOLUTION", "SABZ: signals that see, think, and adapt", tsize=35)
cards = [
    ("SEE", "Cameras count vehicles on every approach, in real time. No new roadworks: SABZ reads the cameras an intersection already has."),
    ("THINK", "A trained DQN brain reads an 18-feature state of the junction and decides, every second, which road gets green."),
    ("ACT", "Signals hand green to the road that needs it, then re-check a second later. One road open at a time."),
]
for i, (h, b) in enumerate(cards):
    x = 0.55 + i * 4.14
    rect(s, x, 2.62, 3.97, 2.0)
    text(s, x + 0.26, 2.84, 3.45, 0.45, [[(h, {"color": GREEN, "font": HEAD, "bold": True, "size": 19})]])
    text(s, x + 0.26, 3.32, 3.48, 1.2, [[(b, {"size": 13.5, "color": RGBColor(0xC7, 0xD2, 0xE8)})]], leading=1.3)
rect(s, 0.55, 4.82, 12.23, 1.1, fill=CARD2, line=GREEN, lw=1.2)
text(s, 0.82, 4.97, 11.7, 0.3, [[("HARD-CODED SAFETY, NEVER LEARNED AWAY", {"color": GREEN, "font": HEAD, "bold": True, "size": 12})]])
text(s, 0.82, 5.28, 11.7, 0.5, [[("Green is never cut under a car inside the 70 m dilemma zone. A starvation guard means no road waits forever. Max-green keeps every split fair.", {"size": 13.5})]])
stats = [("5,506", "model parameters"), ("1 s", "decision cycle"),
         ("2 of 3", "benchmark scenarios won"), ("7.3 s", "best average delay")]
for i, (v, l) in enumerate(stats):
    x = 0.55 + i * 3.11
    rect(s, x, 6.12, 2.94, 0.9)
    text(s, x + 0.2, 6.21, 2.6, 0.45, [[(v, {"color": GREEN, "font": HEAD, "bold": True, "size": 20})]])
    text(s, x + 0.2, 6.66, 2.6, 0.3, [[(l, {"color": MUTED, "size": 10.5})]])
footer(s, 2)

# ============================================================ SLIDE 3: ARCHITECTURE
s = slide(); logo(s, y=0.26, w=1.85)
text(s, 0.55, 1.12, 12.2, 0.5, [[("System Architecture", {"font": HEAD, "bold": True, "size": 27})]])
s.shapes.add_picture("architecture.png", Inches(1.07), Inches(1.68), width=Inches(11.2))
footer(s, 3)

# ============================================================ SLIDE 4: DEMO + RESULTS
s = slide(); logo(s)
kicker_title(s, "LIVE DEMO · RESULTS", "Watch it beat the fixed timer, live", tsize=36)
text(s, 0.55, 2.62, 7.9, 1.05, [[("Two identical intersections. Same rush-hour traffic, same start. One runs a blind 30-second fixed timer. The other runs the real trained SABZ brain, deciding every second, one road open at a time.", {"size": 15, "color": RGBColor(0xC7, 0xD2, 0xE8)})]], leading=1.3)
bullet(s, 0.55, 3.95, 7.9, "Identical traffic generated live on both sides", size=14.5, h=0.4)
bullet(s, 0.55, 4.42, 7.9, "Real trained weights: the same 5,506-parameter brain from our benchmarks", size=14.5, h=0.4)
bullet(s, 0.55, 4.89, 7.9, "Every decision explained on screen, in English and Urdu", size=14.5, h=0.4)
res = [("2/3", "scenarios won against actuated control"),
       ("7.3 s", "best average delay per vehicle"),
       ("47%", "less delay in the best live battle")]
for i, (v, l) in enumerate(res):
    x = 0.55 + i * 2.7
    rect(s, x, 5.5, 2.53, 1.28)
    text(s, x + 0.2, 5.64, 2.15, 0.5, [[(v, {"color": GREEN, "font": HEAD, "bold": True, "size": 23})]])
    text(s, x + 0.2, 6.18, 2.15, 0.55, [[(l, {"color": MUTED, "size": 10.5})]], leading=1.15)
# QR panel
rect(s, 8.85, 2.62, 3.93, 4.16)
rect(s, 9.62, 2.86, 2.4, 2.4, fill=WHITE, line=None, radius=0.045)
s.shapes.add_picture("demo_qr.png", Inches(9.72), Inches(2.96), width=Inches(2.2))
text(s, 8.85, 5.42, 3.93, 0.35, [[("SCAN FOR THE LIVE DEMO", {"color": GREEN, "font": HEAD, "bold": True, "size": 13.5})]], align=PP_ALIGN.CENTER)
text(s, 9.05, 5.8, 3.53, 0.6, [[("Run the battle on your phone: fixed timer vs the SABZ brain.", {"color": MUTED, "size": 11.5})]], align=PP_ALIGN.CENTER, leading=1.25)
text(s, 8.85, 6.42, 3.93, 0.25, [[("github.com/pakmultilinks-dot/sabz", {"color": MUTED, "size": 10})]], align=PP_ALIGN.CENTER)
footer(s, 4)

prs.save("SABZ_Presentation.pptx")
print("saved SABZ_Presentation.pptx, slides:", len(prs.slides.__iter__.__self__._sldIdLst))
